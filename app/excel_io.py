from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import OrganisationResult

OUTPUT_COLUMNS = [
    "NGO Name",
    "City",
    "Role in the deal",
    "Key person",
    "Designation",
    "Email",
    "Confidence",
    "Phone",
    "LinkedIn URL",
    "Alt person",
    "Alt designation",
    "Alt email",
    "Alt phone",
    "Other emails",
    "Other phones / routes",
    "Website",
    "Source",
    "Evidence / relevance",
    "Named programme",
    "What they do",
    "Custom appreciation",
    "Why Xperimentor fits",
    "Notes / next step",
]
INTERNAL_COLUMNS = ["verification_status", "errors"]


def load_workbook_for_enrichment(path: str | Path):
    return load_workbook(path)


def detect_header_row(ws: Worksheet) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
        lowered = {value.lower() for value in values}
        if "ngo name" in lowered and "website" in lowered:
            return row[0].row
    raise ValueError("Could not find a header row containing 'NGO Name' and 'Website'.")


def header_map(ws: Worksheet, header_row: int) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[header_row]
        if cell.value is not None and str(cell.value).strip()
    }


def ensure_columns(ws: Worksheet, header_row: int) -> dict[str, int]:
    headers = header_map(ws, header_row)
    next_col = ws.max_column + 1
    for name in OUTPUT_COLUMNS + INTERNAL_COLUMNS:
        if name not in headers:
            ws.cell(header_row, next_col).value = name
            headers[name] = next_col
            next_col += 1
    return headers


def iter_records(ws: Worksheet, header_row: int, headers: dict[str, int]):
    for row in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row, headers["NGO Name"]).value
        website = ws.cell(row, headers["Website"]).value
        if not name and not website:
            continue
        yield row, str(name or "").strip(), str(website or "").strip()


def is_processed(ws: Worksheet, row: int, headers: dict[str, int], resume: bool) -> bool:
    if not resume:
        return False
    status_col = headers.get("verification_status")
    errors_col = headers.get("errors")
    return bool(
        (status_col and ws.cell(row, status_col).value)
        or (errors_col and ws.cell(row, errors_col).value)
    )


def write_result(ws: Worksheet, row: int, headers: dict[str, int], result: OrganisationResult) -> None:
    values: dict[str, Any] = {
        "City": result.city,
        "Role in the deal": result.role_in_deal,
        "Key person": result.key_person,
        "Designation": result.designation,
        "Email": result.email,
        "Confidence": result.confidence,
        "Phone": result.phone,
        "LinkedIn URL": result.linkedin_url,
        "Alt person": result.alt_person,
        "Alt designation": result.alt_designation,
        "Alt email": result.alt_email,
        "Alt phone": result.alt_phone,
        "Other emails": result.other_emails,
        "Other phones / routes": result.other_phones_routes,
        "Source": result.source,
        "Evidence / relevance": result.evidence_relevance,
        "Named programme": result.named_programme,
        "What they do": result.what_they_do,
        "Custom appreciation": result.custom_appreciation,
        "Why Xperimentor fits": result.why_xperimentor_fits,
        "Notes / next step": result.notes_next_step,
        "verification_status": result.verification_status,
        "errors": "; ".join(sorted(set(result.errors))),
    }
    for name, value in values.items():
        if name in headers and value not in (None, ""):
            ws.cell(row, headers[name]).value = value


def format_output(ws: Worksheet, header_row: int, headers: dict[str, int]) -> None:
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
    ws.auto_filter.ref = ws.dimensions
    low_fill = PatternFill("solid", fgColor="FFF2CC")
    review_fill = PatternFill("solid", fgColor="F4CCCC")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    wrap_cols = {
        "Evidence / relevance",
        "Named programme",
        "What they do",
        "Custom appreciation",
        "Why Xperimentor fits",
        "Notes / next step",
        "Other emails",
        "Other phones / routes",
        "Likely pitch angle",
        "Suggested first email",
        "LinkedIn note",
        "Objection risk",
        "Next action",
    }
    for name, col in headers.items():
        width = min(max(len(name) + 2, 12), 45)
        for row in range(header_row + 1, min(ws.max_row, header_row + 30) + 1):
            value = ws.cell(row, col).value
            if value:
                width = min(max(width, len(str(value)) + 2), 60)
        ws.column_dimensions[get_column_letter(col)].width = width
        if name in wrap_cols:
            for row in range(header_row + 1, ws.max_row + 1):
                ws.cell(row, col).alignment = ws.cell(row, col).alignment.copy(wrap_text=True)
    for row in range(header_row + 1, ws.max_row + 1):
        confidence = ws.cell(row, headers.get("Confidence", 0)).value if headers.get("Confidence") else None
        status = ws.cell(row, headers.get("verification_status", 0)).value if headers.get("verification_status") else ""
        fill = review_fill if status == "REVIEW_REQUIRED" else low_fill if isinstance(confidence, int) and confidence < 60 else None
        if fill:
            for col in range(1, ws.max_column + 1):
                ws.cell(row, col).fill = fill
    for name in ("Website", "LinkedIn URL"):
        col = headers.get(name)
        if col:
            for row in range(header_row + 1, ws.max_row + 1):
                cell = ws.cell(row, col)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
