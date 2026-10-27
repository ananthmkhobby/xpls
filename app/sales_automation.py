from __future__ import annotations

import re
from dataclasses import dataclass
from urllib.parse import quote

from openpyxl.worksheet.worksheet import Worksheet

from app.curated_research import curated_program_research


SALES_COLUMNS = [
    "Lead fit score",
    "Lead priority",
    "Lead segment",
    "Recommended channel",
    "Outreach readiness",
    "Named programme",
    "What they do",
    "Custom appreciation",
    "Why Xperimentor fits",
    "Likely pitch angle",
    "Email subject",
    "Suggested first email",
    "Send email",
    "LinkedIn note",
    "Open LinkedIn",
    "Objection risk",
    "Next action",
]


@dataclass(frozen=True)
class LeadAutomation:
    lead_fit_score: int
    lead_priority: str
    lead_segment: str
    recommended_channel: str
    outreach_readiness: str
    likely_pitch_angle: str
    email_subject: str
    suggested_first_email: str
    linkedin_note: str
    objection_risk: str
    next_action: str


def ensure_sales_columns(ws: Worksheet, header_row: int, headers: dict[str, int]) -> dict[str, int]:
    next_col = ws.max_column + 1
    for name in SALES_COLUMNS:
        if name not in headers:
            ws.cell(header_row, next_col).value = name
            headers[name] = next_col
            next_col += 1
    return headers


def apply_sales_automation(ws: Worksheet, header_row: int, headers: dict[str, int]) -> None:
    ensure_sales_columns(ws, header_row, headers)
    for row in range(header_row + 1, ws.max_row + 1):
        ngo_name = get_cell(ws, row, headers, "NGO Name")
        website = get_cell(ws, row, headers, "Website")
        if not ngo_name and not website:
            continue
        row_data = row_with_curated_research(read_row(ws, row, headers))
        lead = build_lead_automation(row_data)
        values = {
            "Lead fit score": lead.lead_fit_score,
            "Lead priority": lead.lead_priority,
            "Lead segment": lead.lead_segment,
            "Recommended channel": lead.recommended_channel,
            "Outreach readiness": lead.outreach_readiness,
            "Named programme": row_data.get("Named programme", ""),
            "What they do": row_data.get("What they do", ""),
            "Custom appreciation": row_data.get("Custom appreciation", ""),
            "Why Xperimentor fits": row_data.get("Why Xperimentor fits", ""),
            "Likely pitch angle": lead.likely_pitch_angle,
            "Email subject": lead.email_subject,
            "Suggested first email": lead.suggested_first_email,
            "Send email": "",
            "LinkedIn note": lead.linkedin_note,
            "Open LinkedIn": "",
            "Objection risk": lead.objection_risk,
            "Next action": lead.next_action,
        }
        for column, value in values.items():
            ws.cell(row, headers[column]).value = value
        apply_action_links(ws, row, headers, row_data, lead)


def build_lead_automation(row: dict[str, str]) -> LeadAutomation:
    row = row_with_curated_research(row)
    score = lead_fit_score(row)
    priority = priority_for(score, row)
    segment = segment_for(row)
    channel = channel_for(row)
    readiness = readiness_for(row)
    pitch = pitch_angle_for(row, segment)
    objection = objection_risk_for(row)
    next_action = next_action_for(priority, channel, row)
    return LeadAutomation(
        lead_fit_score=score,
        lead_priority=priority,
        lead_segment=segment,
        recommended_channel=channel,
        outreach_readiness=readiness,
        likely_pitch_angle=pitch,
        email_subject=email_subject(row),
        suggested_first_email=email_draft(row, pitch, channel),
        linkedin_note=linkedin_note(row, pitch),
        objection_risk=objection,
        next_action=next_action,
    )


def row_with_curated_research(row: dict[str, str]) -> dict[str, str]:
    curated = curated_program_research(row.get("NGO Name", ""), row.get("Website", ""))
    if not any((curated.named_programme, curated.what_they_do, curated.custom_appreciation, curated.why_xperimentor_fits)):
        return row
    enriched = dict(row)
    defaults = {
        "Named programme": curated.named_programme,
        "What they do": curated.what_they_do,
        "Custom appreciation": curated.custom_appreciation,
        "Why Xperimentor fits": curated.why_xperimentor_fits,
    }
    for key, value in defaults.items():
        if value and not enriched.get(key):
            enriched[key] = value
    return enriched


def lead_fit_score(row: dict[str, str]) -> int:
    text = joined(row, "Role in the deal", "Evidence / relevance", "Notes / next step", "Source")
    score = 0
    scoring_rules = [
        (("stem", "science", "robotics", "coding", "steam"), 25),
        (("school", "student", "classroom", "k-12"), 20),
        (("teacher training", "capacity building", "training"), 15),
        (("government school", "government schools"), 10),
        (("csr", "funder", "sponsor"), 10),
        (("partnership", "partner", "implementation"), 10),
    ]
    for keywords, points in scoring_rules:
        if any(keyword in text for keyword in keywords):
            score += points
    if usable_email(row.get("Email", "")):
        score += 10
    if row.get("Key person") and row.get("Key person", "").lower() not in {"general contact", "not named publicly"}:
        score += 5
    if row.get("LinkedIn URL"):
        score += 5
    if "competitor" in text:
        score -= 15
    if "limited" in text or "no relevant" in text:
        score -= 20
    return max(0, min(100, score))


def priority_for(score: int, row: dict[str, str]) -> str:
    if score >= 80 and "REVIEW_REQUIRED" not in row.get("verification_status", ""):
        return "A - outreach now"
    if score >= 60:
        return "B - warm lead"
    if score >= 40:
        return "C - nurture"
    return "D - manual review"


def segment_for(row: dict[str, str]) -> str:
    text = joined(row, "Role in the deal", "Evidence / relevance", "Notes / next step")
    if any(term in text for term in ("funder", "sponsor", "csr partner", "csr")):
        return "CSR/funding prospect"
    if any(term in text for term in ("implementation", "implementer", "program delivery")):
        return "Implementation partner"
    if any(term in text for term in ("school network", "schools", "classroom")):
        return "School network"
    if any(term in text for term in ("training", "teacher")):
        return "Training partner"
    if "government" in text:
        return "Government-linked partner"
    return "General NGO prospect"


def channel_for(row: dict[str, str]) -> str:
    if usable_email(row.get("Email", "")) and row.get("LinkedIn URL"):
        return "Email + LinkedIn"
    if usable_email(row.get("Email", "")):
        return "Email"
    if row.get("LinkedIn URL"):
        return "LinkedIn"
    if row.get("Phone") and "not" not in row.get("Phone", "").lower():
        return "Phone"
    return "Manual research"


def readiness_for(row: dict[str, str]) -> str:
    if usable_email(row.get("Email", "")) and row.get("Lead priority", "") != "D - manual review":
        return "Ready for outreach"
    if usable_email(row.get("Email", "")):
        return "Verify before outreach"
    if row.get("LinkedIn URL") or row.get("Phone"):
        return "Needs contact verification"
    return "Manual research required"


def pitch_angle_for(row: dict[str, str], segment: str) -> str:
    text = joined(row, "Evidence / relevance", "Notes / next step", "Role in the deal")
    if any(term in text for term in ("stem", "science", "robotics", "coding", "steam")):
        return "Explore a STEM/robotics learning partnership for school students."
    if "teacher training" in text or "training" in text:
        return "Position Xperimentor as a hands-on STEM module for teacher-training programs."
    if "government school" in text or "government schools" in text:
        return "Pitch a government-school STEM implementation pilot with measurable classroom outcomes."
    if "csr" in text or "funder" in text or "sponsor" in text:
        return "Frame the proposal as a CSR-ready STEM education initiative with implementation support."
    if segment == "School network":
        return "Offer a pilot program for selected schools with student engagement metrics."
    return "Start with education-program fit and ask for the right partnership owner."


def objection_risk_for(row: dict[str, str]) -> str:
    text = joined(row, "Notes / next step", "errors", "Evidence / relevance")
    if "competitor" in text:
        return "May be adjacent competitor; position as collaboration or avoid direct product pitch."
    if "no_email_found" in text or "manual review" in text or "blocked" in text:
        return "Contact route uncertain; verify before campaign upload."
    if "limited" in text:
        return "Program fit may be weak; use low-touch nurture."
    return "Normal cold-outreach risk."


def next_action_for(priority: str, channel: str, row: dict[str, str]) -> str:
    if priority.startswith("A") and "Email" in channel:
        return "Send personalized email and follow with LinkedIn touch in 2 days."
    if priority.startswith("B") and "Email" in channel:
        return "Verify contact, then send warm partnership email."
    if channel == "LinkedIn":
        return "Connect with the listed person and ask for the partnership owner."
    if channel == "Phone":
        return "Call office and ask for CSR/program partnerships contact."
    return "Manually verify contact route before outreach."


def email_draft(row: dict[str, str], pitch: str, channel: str) -> str:
    ngo = row.get("NGO Name") or "your organisation"
    appreciation = appreciation_sentence(row)
    programme = programme_reference(row)
    why = row.get("Why Xperimentor fits") or why_reaching_out(row, pitch)
    return (
        f"Dear {ngo} Team,\n\n"
        "I'm Chethan, writing on behalf of Xperimentor "
        "(Scikraft Education and Engineering Design Pvt. Ltd.), a Bengaluru-based organisation "
        "working with corporate CSR teams and NGOs to bring functional, curriculum-aligned Physics, "
        "Chemistry, Biology, and Mathematics labs into government school classrooms.\n\n"
        f"We admire {ngo}'s work. {appreciation}\n\n"
        "What we bring:\n\n"
        "Our Lab-in-Classroom system turns any existing classroom into a functional science lab, "
        "with no separate room or extra infrastructure. It includes:\n\n"
        "- 150+ hands-on experiments across Physics, Chemistry, Biology, and Mathematics, "
        "aligned to CBSE, ICSE, and State Board curricula\n"
        "- Structured teacher training and mentoring, with onboarding, classroom workshops, "
        "and year-round support\n"
        "- AI-powered assessment and impact tracking, with real-time dashboards and CSR-ready "
        "reports through our analytics layer, TiQer\n\n"
        "Our current scale:\n\n"
        "- 100+ government schools reached via CSR across 4 states\n"
        "- 10,000+ students impacted, 500+ teachers trained\n"
        "- 94.9% active teacher adoption rate\n\n"
        "Why we're reaching out:\n\n"
        f"{why} {programme}\n\n"
        "Could you direct this to your Corporate Partnerships or Programmes team? We'd welcome "
        "a short call, and we're happy to share our brochure and case studies.\n\n"
        "Best,\n"
        "Chethan"
    )


def email_subject(row: dict[str, str]) -> str:
    ngo = row.get("NGO Name") or "your education work"
    text = joined(row, "Named programme", "What they do", "Evidence / relevance", "Notes / next step", "Role in the deal")
    if any(term in text for term in ("stem", "science", "robotics", "coding", "steam")):
        theme = "Science and Math Lab-in-Classroom Kits"
    elif "teacher" in text or "training" in text:
        theme = "Hands-on STEM Labs for Teacher-Led Classrooms"
    elif "csr" in text or "funder" in text or "sponsor" in text:
        theme = "CSR-Ready Science and Math Lab Programme"
    else:
        theme = "Hands-on Science and Math Learning"
    return f"Partnership Opportunity: {theme} for {ngo}"


def mailto_url(row: dict[str, str], lead: LeadAutomation) -> str:
    email = row.get("Email", "")
    if not usable_email(email):
        return ""
    body = lead.suggested_first_email[:1900]
    return f"mailto:{quote(email)}?subject={quote(lead.email_subject)}&body={quote(body)}"


def apply_action_links(ws: Worksheet, row_number: int, headers: dict[str, int], row: dict[str, str], lead: LeadAutomation) -> None:
    send_col = headers.get("Send email")
    if send_col:
        send_cell = ws.cell(row_number, send_col)
        url = mailto_url(row, lead)
        if url:
            send_cell.value = "Send email"
            send_cell.hyperlink = url
            send_cell.style = "Hyperlink"
        else:
            send_cell.value = "Verify contact"

    linkedin_col = headers.get("Open LinkedIn")
    if linkedin_col:
        linkedin_cell = ws.cell(row_number, linkedin_col)
        linkedin_url = row.get("LinkedIn URL", "")
        if linkedin_url.startswith("http"):
            linkedin_cell.value = "Open LinkedIn"
            linkedin_cell.hyperlink = linkedin_url
            linkedin_cell.style = "Hyperlink"
        else:
            linkedin_cell.value = ""


def linkedin_note(row: dict[str, str], pitch: str) -> str:
    ngo = row.get("NGO Name") or "your organisation"
    compact_pitch = re.sub(r"\s+", " ", pitch).rstrip(".")
    return f"Hi, I came across {ngo}'s education work. Would love to connect and explore whether {compact_pitch.lower()} could be relevant."


def create_sales_sheets(wb, source_ws: Worksheet, header_row: int, headers: dict[str, int]) -> None:
    for sheet_name in ("Priority Outreach", "Manual Review"):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
    priority_ws = wb.create_sheet("Priority Outreach")
    review_ws = wb.create_sheet("Manual Review")
    selected = [
        "NGO Name",
        "Lead fit score",
        "Lead priority",
        "Lead segment",
        "Recommended channel",
        "Named programme",
        "What they do",
        "Custom appreciation",
        "Why Xperimentor fits",
        "Key person",
        "Designation",
        "Email",
        "Send email",
        "Email subject",
        "LinkedIn URL",
        "Open LinkedIn",
        "Likely pitch angle",
        "Suggested first email",
        "Next action",
        "Website",
    ]
    write_sheet_header(priority_ws, selected)
    write_sheet_header(review_ws, selected + ["errors"])
    priority_row = 2
    review_row = 2
    for row in range(header_row + 1, source_ws.max_row + 1):
        values = {column: get_cell(source_ws, row, headers, column) for column in selected + ["errors"]}
        if not values.get("NGO Name"):
            continue
        if str(values.get("Lead priority", "")).startswith(("A", "B")):
            write_values(priority_ws, priority_row, selected, values)
            apply_action_links(priority_ws, priority_row, header_lookup(selected), values, build_lead_automation(values))
            priority_row += 1
        if values.get("errors") or values.get("Recommended channel") == "Manual research":
            write_values(review_ws, review_row, selected + ["errors"], values)
            apply_action_links(review_ws, review_row, header_lookup(selected + ["errors"]), values, build_lead_automation(values))
            review_row += 1
    autosize(priority_ws)
    autosize(review_ws)


def read_row(ws: Worksheet, row: int, headers: dict[str, int]) -> dict[str, str]:
    return {name: get_cell(ws, row, headers, name) for name in headers}


def get_cell(ws: Worksheet, row: int, headers: dict[str, int], name: str) -> str:
    col = headers.get(name)
    if not col:
        return ""
    value = ws.cell(row, col).value
    return str(value).strip() if value is not None else ""


def joined(row: dict[str, str], *fields: str) -> str:
    return " ".join(row.get(field, "") for field in fields).lower()


def usable_email(value: str) -> bool:
    lowered = (value or "").lower()
    return bool(value and "@" in value and "not " not in lowered and "example." not in lowered and "test." not in lowered)


def first_name(name: str) -> str:
    if not name or name.lower() in {"general contact", "not named publicly"}:
        return "there"
    return name.split()[0].strip(",")


def appreciation_sentence(row: dict[str, str]) -> str:
    custom = row.get("Custom appreciation", "")
    if custom:
        return custom
    evidence = clean_generated_sentence(row.get("Evidence / relevance", ""))
    named = row.get("Named programme", "")
    what = row.get("What they do", "")
    notes = clean_generated_sentence(row.get("Notes / next step", ""))
    text = joined(row, "Named programme", "What they do", "Evidence / relevance", "Notes / next step", "Role in the deal")
    if named and what:
        first_program = named.split(";", 1)[0]
        detail = remove_repeated_heading(what, first_program)
        return (
            f"We were glad to see your work through {first_program}, particularly its focus on "
            f"{lower_first(detail)}"
        )
    if named:
        return f"We were glad to see initiatives such as {named}."
    if evidence and not evidence.lower().startswith("limited"):
        return (
            f"We were glad to see your focus on {human_program_theme(text)}, "
            f"especially the public work described in your materials: {evidence}"
        )
    if notes and "limited" not in notes.lower():
        return f"We were glad to learn about your education and community programmes: {notes}"
    return "We were glad to learn about your public work with communities, schools, and young people."


def programme_reference(row: dict[str, str]) -> str:
    fit = row.get("Why Xperimentor fits", "")
    if fit:
        return "This can be shaped as a practical pilot or CSR-backed implementation with clear classroom adoption and impact reporting."
    segment = segment_for(row)
    if segment == "CSR/funding prospect":
        return "This could give your CSR partners a measurable academic outcome to report alongside their existing education or infrastructure work."
    if segment == "Implementation partner":
        return "This could add a practical science and math layer to programmes you already implement in schools."
    if segment == "Training partner":
        return "This could strengthen teacher-training outcomes by giving teachers ready-to-use experiments and classroom mentoring."
    if segment == "School network":
        return "This could help selected schools add hands-on science and math learning without waiting for a separate lab room."
    return "This could create a measurable, classroom-level STEM outcome within your existing education initiatives."


def why_reaching_out(row: dict[str, str], pitch: str) -> str:
    text = joined(row, "Named programme", "What they do", "Evidence / relevance", "Notes / next step", "Role in the deal")
    if "government school" in text or "government schools" in text:
        return "Given your work with government schools, we see potential to add hands-on science and math learning to classrooms you already support."
    if "teacher" in text or "training" in text:
        return "Given your focus on teacher capacity and classroom practice, we see a strong fit for a practical Lab-in-Classroom model."
    if "csr" in text or "funder" in text or "sponsor" in text:
        return "Given your CSR and partnership orientation, we see potential for a structured STEM education programme that is simple to deploy and easy to report."
    if "school" in text or "student" in text:
        return "Given your reach with schools and students, we see potential to add a hands-on science and math component to the learning environment."
    return pitch.rstrip(".") + "."


def human_program_theme(text: str) -> str:
    if any(term in text for term in ("stem", "science", "robotics", "coding", "steam")):
        return "STEM and science education"
    if "teacher" in text or "training" in text:
        return "teacher training and classroom learning"
    if "government school" in text or "government schools" in text:
        return "government school education"
    if "csr" in text:
        return "CSR-backed education initiatives"
    if "school" in text:
        return "school education"
    return "education and youth development"


def lower_first(text: str) -> str:
    if not text:
        return text
    return text[0].lower() + text[1:]


def remove_repeated_heading(text: str, heading: str) -> str:
    cleaned = re.sub(r"\s+", " ", text or "").strip()
    pattern = re.escape(heading).replace("\\ ", r"\s+")
    cleaned = re.sub(rf"^{pattern}\s*[:\-–—]\s*", "", cleaned, flags=re.I)
    return cleaned or text


def clean_generated_sentence(text: str) -> str:
    sentence = short_sentence(text)
    replacements = {
        "Found education/STEM signals:": "education/STEM signals include",
        "SUBJECT MATCH —": "",
        "CLOSEST SUBJECT MATCH ON THE LIST —": "",
        "TIGHTEST MISSION MATCH ON THE ENTIRE LIST —": "",
    }
    for old, new in replacements.items():
        sentence = sentence.replace(old, new)
    return sentence.strip()


def short_sentence(text: str) -> str:
    if not text:
        return "Your public work appears relevant to education outreach."
    sentence = re.split(r"(?<=[.!?])\s+", text.strip())[0]
    return sentence[:180].rstrip(" ,;") + ("." if not sentence.endswith((".", "!", "?")) else "")


def write_sheet_header(ws: Worksheet, columns: list[str]) -> None:
    for index, column in enumerate(columns, start=1):
        ws.cell(1, index).value = column
        ws.cell(1, index).font = ws.cell(1, index).font.copy(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_values(ws: Worksheet, row: int, columns: list[str], values: dict[str, str]) -> None:
    for index, column in enumerate(columns, start=1):
        ws.cell(row, index).value = values.get(column, "")


def header_lookup(columns: list[str]) -> dict[str, int]:
    return {column: index for index, column in enumerate(columns, start=1)}


def autosize(ws: Worksheet) -> None:
    for column_cells in ws.columns:
        width = 12
        for cell in column_cells[:30]:
            if cell.value:
                width = min(max(width, len(str(cell.value)) + 2), 60)
        ws.column_dimensions[column_cells[0].column_letter].width = width
